# -*- coding: utf-8 -*-
"""Convert a flat Jasper export (.md) to a readable .docx,
styled like 'TotalRecalls Video Storyboard.md.docx'.

Importable:  from convert import convert_md
             convert_md(src_md, out_docx) -> dict(counts)
CLI:         python convert.py [SRC.md [OUT.docx]]
"""
import os
import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
TITLE_PT, MAJ_PT, SUB_PT, BODY_PT = 24, 18, 13.5, 12
AFTER = Pt(5)

REF = r"C:\Users\break\Projects\TotalRecalls\Marketing\Jasper\YouTube\TotalRecalls Video Storyboard.md.docx"
DEFAULT_SRC = r"C:\Users\break\Projects\TotalRecalls\Marketing\Jasper\Jasper Plan 1.md"
DEFAULT_OUT = r"C:\Users\break\Projects\TotalRecalls\Marketing\Jasper\Jasper Plan 1.docx"

def _base_doc(ref):
    """Template doc: the storyboard file if present, else a blank one."""
    if ref and os.path.exists(ref):
        return Document(ref)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(72)  # 1 inch
    return doc

# ------------------------------------------------- stateless helpers
UI_ARTIFACT = {"Create Voice", "Share", "Content generated", "Untitled Document"}
def is_artifact(t):
    t = t.strip()
    return t in UI_ARTIFACT or t.lower().startswith("jasper tip")

LABEL  = re.compile(r"^([A-Za-z][A-Za-z0-9 &/.()-]{0,19}):\s+(\S.*)$")
def label_parts(t):
    m = LABEL.match(t)
    if not m: return None
    lab = m.group(1)
    if len(lab.split()) > 3: return None          # not a label, a sentence fragment
    return (lab + ": ", m.group(2))

TRAIL = re.compile(r"^([A-Za-z][A-Za-z0-9 &/.()-]{0,19}):$")
def is_trail(t): return bool(TRAIL.match(t.strip()))

NUMITEM = re.compile(r"^(Email \d|Week \d|Day \d|Scene \d|\[\d|Video \d|Step \d|(\d+)\.|(Meta|LinkedIn|X) Ad \d)")
def is_numbered(t): return bool(NUMITEM.match(t))

def heading_like(t):
    t = t.strip()
    if not (8 <= len(t) <= 45) or not re.search(r"[A-Za-z]", t): return False
    if t[-1] in ".?!:,": return False
    if t[0] in "[(#\u2014-": return False
    if is_artifact(t) or label_parts(t) or is_trail(t): return False
    if " " not in t: return False              # single word -> body (avoids 'Status','KPI','Blog')
    if not (t[0].isupper() or t[0].isdigit()): return False
    if "," in t and not (re.match(r"^\d{1,2}[.)]\s+\S", t) and len(t) <= 45):
        return False                            # list fragment (numbered sections keep the comma)
    return True

def header_like(s):
    s = s.strip()
    if not (1 <= len(s) <= 30): return False
    if s[-1] in ".!?:": return False
    if re.fullmatch(r"[\d⬜_\-x/ ]+", s): return False
    return True

def is_key(s):
    s = s.strip()
    return 1 <= len(s) <= 32

HDR_TOKENS = {"CTA","KPI","Owner","Channel","Asset","Status","Dependencies","Day",
              "Date","Platform","Content Type","Topic","Description","Audience",
              "Window","Action","Rule","Logic","Bucket","Primary CTA","Receives",
              "Stage","Trigger","Definition","Due Date","Section","Duration",
              "Key Visual","Audio Tone","How","When","Success metric","Element",
              "What to test","Role","Focus","Cadence","Week","Platform(s)"}

MAJOR_KW = re.compile(
    r"(TotalRecalls|Email Sequence|Email Campaign|Search Ads|Meta Ads|LinkedIn Ads|X Ads|"
    r"Landing Page|FAQ Page|Pricing Page|Guide Pages|Content Calendar|Video Scripts|"
    r"Timeline|Deliverables|Objectives|Key Messages|Channel Mix|Posting Cadence|"
    r"Short-Form Video|Social Media Posts|Ad Campaign|Welcome Email|Upgrade Email|"
    r"Consolidation Email|Ownership Email|Search Campaign|Google Ads|Facebook Post|LinkedIn Post|"
    r"Instagram Post|X Ad|Video Storyboard|Landing|Strategy Framework|Multi-AI|Privacy Features|"
    r"Persuasion Email|Reminder Email|Discount Reminder|Final Reminder|Publishing schedule|"
    r"Summary Table|Content Plan)")
def is_major(t):
    return bool(MAJOR_KW.search(t)) and heading_like(t)

# ------------------------------------------------- table detection
def try_table(items, prev_last):
    """items: run cells. prev_last: last line of preceding block (or None).
    Returns (n, rows, pattern, recovered) or None."""
    m = len(items)
    if m < 8: return None
    cands = []
    for n in range(2, 9):
        for pattern in ("A", "B"):
            off = n if pattern == "A" else (n - 1)
            if off < 1 or m < off: continue
            if (m - off) % n != 0: continue
            rows = (m - off) // n
            if rows < 3: continue
            hdrs = list(items[:off])
            recovered = None
            if pattern == "B":
                if prev_last is None or not header_like(prev_last): continue
                recovered = prev_last
                hdrs = [recovered] + hdrs
            if not all(header_like(h) for h in hdrs): continue
            data = [items[off + r * n: off + (r + 1) * n] for r in range(rows)]
            c0 = [row[0] for row in data]
            if not all(is_key(c) for c in c0): continue
            a0 = sum(len(c) for c in c0) / rows
            oth = [x for row in data for x in row[1:]]
            ao = sum(len(x) for x in oth) / len(oth)
            if ao < a0 * 1.3: continue
            if not any(len(x) >= 12 for x in oth): continue   # run must carry real content somewhere
            # reject if the first data row is really the tail of a longer header
            if sum(1 for x in data[0] if x.strip() in HDR_TOKENS) >= n - 1: continue
            cands.append((n, rows, pattern, recovered))
    if not cands: return None
    return max(cands, key=lambda c: (c[1], c[0]))   # prefer most rows; tie -> more columns

# ------------------------------------------------- pipeline
def convert_md(src, out, dbg=None, ref=REF):
    raw = open(src, encoding="utf-8").read()
    lines = raw.split("\n")
    blocks = []
    cur, start = [], 0
    for i, ln in enumerate(lines):
        if ln.strip() == "":
            if cur:
                blocks.append({"start": start, "lines": cur}); cur = []
        else:
            if not cur: start = i
            cur.append(ln)
    if cur: blocks.append({"start": start, "lines": cur})

    # table detection
    table_cells = set()
    tables = {}
    consumed_prev_last = set()
    i = 0
    while i < len(blocks):
        if len(blocks[i]["lines"]) != 1:
            i += 1; continue
        j = i
        while j < len(blocks) and len(blocks[j]["lines"]) == 1:
            j += 1
        items = [blocks[k]["lines"][0] for k in range(i, j)]
        prev_last = None
        if i > 0:
            pb = blocks[i - 1]["lines"]
            prev_last = pb[-1] if len(pb) >= 2 else None
        d = try_table(items, prev_last)
        if d:
            tables[i] = d
            for k in range(i, j): table_cells.add(k)
            if d[2] == "B":
                consumed_prev_last.add(i - 1)
        i = j

    first_major_idx = None
    for i, b in enumerate(blocks):
        if i in table_cells: continue
        if len(b["lines"]) == 1 and is_major(b["lines"][0]):
            first_major_idx = i
            break

    # doc build
    doc = _base_doc(ref)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"): body.remove(child)

    def _font(r, size, bold=None, italic=None):
        r.font.name = FONT; r.font.size = Pt(size)
        if bold is not None: r.font.bold = bold
        if italic is not None: r.font.italic = italic
        return r

    def para(spacing_after=AFTER):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = spacing_after
        p.paragraph_format.line_spacing = 1.0
        return p

    def add_title(t):  p = para(); _font(p.add_run(t), TITLE_PT, bold=True)
    def add_major(t):  p = para(); _font(p.add_run(t), MAJ_PT, bold=True)
    def add_sub(t):    p = para(); _font(p.add_run(t), SUB_PT, bold=True)
    def add_note(t):   p = para(); _font(p.add_run(t), 10, italic=True)

    def add_body(t):
        p = para()
        lp = label_parts(t)
        if lp:
            _font(p.add_run(lp[0]), BODY_PT, bold=True)
            _font(p.add_run(lp[1]), BODY_PT)
        elif is_trail(t):
            _font(p.add_run(t), BODY_PT, bold=True)
        else:
            _font(p.add_run(t), BODY_PT)
        return p

    def classify(t):
        if is_artifact(t): return ("note", t)
        if is_major(t):    return ("major", t)
        if is_numbered(t): return ("sub", t)
        if heading_like(t):return ("sub", t)
        return ("body", t)

    def ensure_grid_style():
        try:
            return doc.styles["Table Grid"]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            st = doc.styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
                el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
                borders.append(el)
            tblPr = st.element.find(qn("w:tblPr"))
            if tblPr is None:
                tblPr = OxmlElement("w:tblPr"); st.element.append(tblPr)
            tblPr.append(borders)
            return st

    def add_table(n, rows, headers):
        t = doc.add_table(rows=1 + len(rows), cols=n)
        t.style = ensure_grid_style()
        from docx.enum.table import WD_TABLE_ALIGNMENT
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        from docx.shared import Inches
        usable = Inches(6.5)
        def colmax(ci):
            v = len(str(headers[ci])) if ci < len(headers) else 0
            for row in rows:
                if ci < len(row): v = max(v, len(str(row[ci])))
            return v
        lens = [colmax(ci) for ci in range(n)]
        tot = sum(lens)
        mins = [0.06] * n
        w = [max(lens[ci] / tot, mins[ci]) for ci in range(n)]
        sw = sum(w)
        w = [usable * (x / sw) for x in w]
        for ci in range(n):
            for ri in range(len(rows) + 1):
                t.cell(ri, ci).width = w[ci]
        for ci, h in enumerate(headers):
            c = t.cell(0, ci); c.paragraphs[0].clear()
            _font(c.paragraphs[0].add_run(h), BODY_PT, bold=True)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                c = t.cell(ri + 1, ci); c.paragraphs[0].clear()
                _font(c.paragraphs[0].add_run(val), BODY_PT)
        para(spacing_after=Pt(4))

    # render
    i = 0
    title_done = False
    last_rendered = None            # de-dupe consecutive identical lines (Jasper repeats titles)
    counts = {"title":0,"major":0,"sub":0,"body":0,"note":0,"table":0}
    report = []
    while i < len(blocks):
        if i in tables:
            n, rows, pattern, recovered = tables[i]
            off = n if pattern == "A" else (n - 1)
            total = off + rows * n
            items = [blocks[i + k]["lines"][0] for k in range(total)]
            data = [items[off + r * n: off + (r + 1) * n] for r in range(rows)]
            headers = ([recovered] if pattern == "B" else []) + list(items[:off])
            add_table(n, data, headers)
            counts["table"] += 1
            report.append(f"table @ line {blocks[i]['start']}: {n} cols x {rows} rows pattern={pattern} "
                          f"headers={headers}\n  row1={data[0] if data else []}")
            i += total
            continue
        b = blocks[i]
        consumed = i in consumed_prev_last
        if consumed:
            L = b["lines"][:-1]
        else:
            L = list(b["lines"])
        for ln in L:
            if ln.strip() == last_rendered:
                continue            # Jasper repeats a line verbatim (e.g. the doc title)
            last_rendered = ln.strip()
            if i not in table_cells:
                if not title_done and i == first_major_idx and ln == b["lines"][0]:
                    add_title(ln); title_done = True; counts["title"] += 1
                    continue
            kind, text = classify(ln)
            if i == first_major_idx and ln == b["lines"][0] and not title_done and kind == "major":
                add_title(ln); title_done = True; counts["title"] += 1
                continue
            {"major": add_major, "sub": add_sub, "note": add_note, "body": add_body, "title": add_title}[kind](text)
            counts[kind] += 1
        i += 1

    doc.save(out)
    if dbg:
        open(dbg, "w", encoding="utf-8").write("\n".join(report))
    return counts

if __name__ == "__main__":
    import sys, os
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
    out = sys.argv[2] if len(sys.argv) > 2 else (os.path.splitext(src)[0] + ".docx")
    print("saved:", out, convert_md(src, out))
